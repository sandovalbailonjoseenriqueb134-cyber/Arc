import os
import pandas as pd
from openpyxl.chart import BarChart, PieChart, Reference
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


def generar_reporte_excel(
    resumen_capas,
    clasificacion_ia,
    catalogo_precios=None,
    nombre_archivo='Reporte_Cuantificacion_Arc.xlsx',
):
  """Genera el reporte en Excel (.xlsx) excluyendo el ruido CAD de los totales principales."""
  print(
      f'\n📊 Arc Exporter: Filtrando ruido CAD y generando reporte final'
      f' ({nombre_archivo})...'
  )

  if catalogo_precios is None:
    catalogo_precios = {
        'Cimentación': 2800.0,
        'Columnas': 3200.0,
        'Vigas/Trabes': 3000.0,
        'Muros': 450.0,
        'Acero/Refuerzo': 35.0,
        'Anotaciones/Detalles': 0.0,
        'Desconocido': 0.0,
    }

  mapa_ia = {}
  if clasificacion_ia and isinstance(clasificacion_ia, list):
    for item in clasificacion_ia:
      mapa_ia[item.get('capa')] = item

  filas = []
  for capa, datos_geom in resumen_capas.items():
    info_ia = mapa_ia.get(
        capa,
        {
            'categoria': 'Sin clasificar',
            'unidad_medida': 'm2',
            'peralte_m': 0.0,
            'es_ruido': False,
            'explicacion': 'Sin información de IA',
        },
    )

    cat = info_ia.get('categoria', 'Sin clasificar')
    unidad = str(info_ia.get('unidad_medida', 'm2')).lower()
    es_ruido = bool(info_ia.get('es_ruido', False))

    area_m2 = datos_geom.get('area_total_m2', 0.0)
    perimetro_m = datos_geom.get('perimetro_total_m', 0.0)
    cant_piezas = datos_geom.get('cantidad_elementos', 0)
    peralte_m = float(info_ia.get('peralte_m') or 0.0)
    volumen_m3 = round(area_m2 * peralte_m, 4)

    # Si es clasificado como ruido o anotación, su costo directo es $0.00
    if es_ruido or cat == 'Anotaciones/Detalles':
      precio_u = 0.0
      costo_total = 0.0
    else:
      if 'm3' in unidad:
        cant_base = volumen_m3
      elif 'm2' in unidad:
        cant_base = area_m2
      elif 'm' in unidad:
        cant_base = perimetro_m
      else:
        cant_base = float(cant_piezas)

      precio_u = float(catalogo_precios.get(cat, 0.0))
      costo_total = round(cant_base * precio_u, 2)

    filas.append({
        'Capa CAD': capa,
        'Categoría Estructural': cat,
        '¿Es Ruido CAD?': 'SÍ (Excluido)' if es_ruido else 'NO',
        'Unidad de Medida': info_ia.get('unidad_medida', 'm2'),
        'Cantidad Elementos': cant_piezas,
        'Área Total (m²)': area_m2,
        'Perímetro / Longitud (m)': perimetro_m,
        'Peralte / Altura (m)': peralte_m,
        'Volumen Total (m³)': volumen_m3,
        'Precio Unitario ($)': precio_u,
        'Costo Total ($)': costo_total,
        'Justificación / Criterio IA': info_ia.get('explicacion', ''),
    })

  df_detalle = pd.DataFrame(filas)

  # Filtrar de la tabla de Resumen por Categoría las capas que sean ruido
  df_util = df_detalle[df_detalle['¿Es Ruido CAD?'] == 'NO']

  df_resumen_cat = (
      df_util.groupby('Categoría Estructural')
      .agg({
          'Cantidad Elementos': 'sum',
          'Área Total (m²)': 'sum',
          'Perímetro / Longitud (m)': 'sum',
          'Volumen Total (m³)': 'sum',
          'Costo Total ($)': 'sum',
      })
      .reset_index()
  )

  ruta_salida = os.path.join(os.getcwd(), nombre_archivo)

  try:
    with pd.ExcelWriter(ruta_salida, engine='openpyxl') as writer:
      df_resumen_cat.to_excel(
          writer, sheet_name='Resumen por Categoría', index=False
      )
      df_detalle.to_excel(writer, sheet_name='Detalle por Capa', index=False)

      azul_encabezado = PatternFill(
          start_color='1F4E78', end_color='1F4E78', fill_type='solid'
      )
      gris_cebra = PatternFill(
          start_color='F9FAFB', end_color='F9FAFB', fill_type='solid'
      )
      fuente_encabezado = Font(
          name='Calibri', size=11, bold=True, color='FFFFFF'
      )
      fuente_datos = Font(name='Calibri', size=11)
      fuente_total = Font(name='Calibri', size=11, bold=True)

      borde_fino = Side(style='thin', color='D9D9D9')
      borde_celda = Border(
          left=borde_fino, right=borde_fino, top=borde_fino, bottom=borde_fino
      )
      borde_total = Border(
          top=Side(style='thin', color='000000'),
          bottom=Side(style='double', color='000000'),
      )

      for sheet_name in writer.sheets:
        ws = writer.sheets[sheet_name]
        ws.views.sheetView[0].showGridLines = True

        max_row = ws.max_row
        max_col = ws.max_column

        ws.row_dimensions[1].height = 28
        for col in range(1, max_col + 1):
          cell = ws.cell(row=1, column=col)
          cell.fill = azul_encabezado
          cell.font = fuente_encabezado
          cell.alignment = Alignment(
              horizontal='center', vertical='center', wrap_text=True
          )

        for row in range(2, max_row + 1):
          ws.row_dimensions[row].height = 20
          es_par = row % 2 == 0

          for col in range(1, max_col + 1):
            cell = ws.cell(row=row, column=col)
            cell.font = fuente_datos
            cell.border = borde_celda
            if es_par:
              cell.fill = gris_cebra

            val = cell.value
            header_name = str(ws.cell(row=1, column=col).value or '')

            if isinstance(val, (int, float)):
              if 'Cantidad' in header_name:
                cell.number_format = '#,##0'
              elif 'Peralte' in header_name:
                cell.number_format = '0.00'
              elif 'Precio' in header_name or 'Costo' in header_name:
                cell.number_format = '"$"#,##0.00'
              else:
                cell.number_format = '#,##0.00'
              cell.alignment = Alignment(horizontal='right', vertical='center')
            elif col in [1, 2, 3, 4]:
              cell.alignment = Alignment(horizontal='left', vertical='center')
            else:
              cell.alignment = Alignment(
                  horizontal='left', vertical='center', wrap_text=True
              )

        # Totales
        fila_totales = max_row + 1
        ws.row_dimensions[fila_totales].height = 24
        cell_label = ws.cell(row=fila_totales, column=1)
        cell_label.value = 'TOTAL'
        cell_label.font = fuente_total
        cell_label.border = borde_total
        cell_label.alignment = Alignment(horizontal='left', vertical='center')

        for col in range(2, max_col + 1):
          cell = ws.cell(row=fila_totales, column=col)
          cell.font = fuente_total
          cell.border = borde_total

          header_name = str(ws.cell(row=1, column=col).value or '')
          col_letter = get_column_letter(col)

          if any(
              k in header_name
              for k in ['Cantidad', 'Área', 'Perímetro', 'Volumen', 'Costo']
          ):
            cell.value = f'=SUM({col_letter}2:{col_letter}{max_row})'
            if 'Costo' in header_name:
              cell.number_format = '"$"#,##0.00'
            elif 'Cantidad' in header_name:
              cell.number_format = '#,##0'
            else:
              cell.number_format = '#,##0.00'
            cell.alignment = Alignment(horizontal='right', vertical='center')

        # Autoajuste
        for col in range(1, max_col + 1):
          col_letter = get_column_letter(col)
          header_name = str(ws.cell(row=1, column=col).value or '')
          if 'Justificación' in header_name:
            ws.column_dimensions[col_letter].width = 50
          else:
            max_len = max(
                len(str(ws.cell(row=r, column=col).value or ''))
                for r in range(1, max_row + 1)
            )
            ws.column_dimensions[col_letter].width = max(max_len + 4, 14)

      # --- DASHBOARD GRÁFICO (SOLO CON DATOS CONSTRUCTIVOS) ---
      ws_resumen = writer.sheets['Resumen por Categoría']
      num_filas_cat = len(df_resumen_cat)

      if num_filas_cat > 0:
        labels = Reference(
            ws_resumen, min_col=1, min_row=2, max_row=num_filas_cat + 1
        )
        datos_costo = Reference(
            ws_resumen, min_col=6, min_row=1, max_row=num_filas_cat + 1
        )

        pie = PieChart()
        pie.add_data(datos_costo, titles_from_data=True)
        pie.set_categories(labels)
        pie.title = 'Distribución del Presupuesto Directo ($)'
        pie.width = 16
        pie.height = 11

        datos_volumen = Reference(
            ws_resumen, min_col=5, min_row=1, max_row=num_filas_cat + 1
        )

        bar = BarChart()
        bar.type = 'col'
        bar.style = 10
        bar.title = 'Volumen Concreto por Categoría (m³)'
        bar.y_axis.title = 'Volumen (m³)'
        bar.x_axis.title = 'Categoría'
        bar.add_data(datos_volumen, titles_from_data=True)
        bar.set_categories(labels)
        bar.legend = None
        bar.width = 16
        bar.height = 11

        ws_resumen.add_chart(pie, 'H2')
        ws_resumen.add_chart(bar, 'H19')

    print(
        f'   ✔ Reporte auditado y sin ruido completado en:\n     ->'
        f' {ruta_salida}\n'
    )
    return ruta_salida

  except Exception as e:
    print(f'   [!] Error al generar Excel: {e}')
    return None